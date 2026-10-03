import json
import os
import re
import sys
from pathlib import Path

import requests
from openai import OpenAI
from pptx import Presentation
from pptx.util import Inches, Pt

client = OpenAI(api_key=os.environ.get("OPENAI_API_KEY"))


def clean_json(text: str) -> str:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    return text.strip()


def validate_slide(slide: dict) -> None:
    required = ("title", "bullets", "image_prompt")
    if any(not slide.get(key) for key in required):
        raise ValueError(f"Incomplete slide record: {slide!r}")
    if not isinstance(slide["bullets"], list) or not all(
        isinstance(item, str) and item.strip() for item in slide["bullets"]
    ):
        raise ValueError("bullets must be a nonempty list of strings")


def generate_slides(topic: str) -> list[dict]:
    response = client.chat.completions.create(
        model=os.environ.get("OPENAI_TEXT_MODEL", "gpt-4o-mini"),
        temperature=0.3,
        messages=[
            {
                "role": "system",
                "content": (
                    "Return JSON only: an object with a slides array containing "
                    "exactly four records. Each record has title, bullets, and "
                    "image_prompt. bullets is an array of concise strings."
                ),
            },
            {
                "role": "user",
                "content": f"Create an educational presentation about: {topic}",
            },
        ],
    )
    payload = json.loads(clean_json(response.choices[0].message.content))
    slides = payload.get("slides")
    if not isinstance(slides, list) or len(slides) != 4:
        raise ValueError("Expected exactly four slide records")
    for slide in slides:
        validate_slide(slide)
    return slides


def download_image(prompt: str, destination: Path) -> None:
    response = client.images.generate(
        model=os.environ.get("OPENAI_IMAGE_MODEL", "dall-e-3"),
        prompt=prompt,
        size="1024x1024",
        n=1,
    )
    url = response.data[0].url if response.data else None
    if not url:
        raise RuntimeError("Image response contained no URL")
    image = requests.get(url, timeout=60)
    image.raise_for_status()
    destination.write_bytes(image.content)
    if destination.stat().st_size == 0:
        raise RuntimeError(f"Empty image asset: {destination}")


def build_deck(generated: list[tuple[dict, Path]], output_path: Path) -> None:
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]
    for slide_data, image_path in generated:
        slide = presentation.slides.add_slide(blank_layout)
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(5.5), Inches(0.7))
        title_frame = title_box.text_frame
        title_frame.text = slide_data["title"]
        title_frame.paragraphs[0].font.size = Pt(26)
        body_box = slide.shapes.add_textbox(Inches(0.65), Inches(1.3), Inches(5.0), Inches(4.8))
        body_frame = body_box.text_frame
        body_frame.text = "\n".join(f"• {bullet}" for bullet in slide_data["bullets"])
        for paragraph in body_frame.paragraphs:
            paragraph.font.size = Pt(18)
        slide.shapes.add_picture(str(image_path), Inches(6.2), Inches(1.0), width=Inches(6.4))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)


def run_pipeline(topic: str, output_path: Path) -> Path:
    slides = generate_slides(topic)
    asset_dir = output_path.parent / "assets"
    asset_dir.mkdir(parents=True, exist_ok=True)
    generated = []
    for number, slide in enumerate(slides, start=1):
        image_path = asset_dir / f"slide_{number:02d}.png"
        print(f"Generating image {number}/{len(slides)}")
        download_image(slide["image_prompt"], image_path)
        generated.append((slide, image_path))
    build_deck(generated, output_path)
    return output_path


if __name__ == "__main__":
    topic = " ".join(sys.argv[1:]).strip() or "Solar energy for small businesses"
    print(f"Created {run_pipeline(topic, Path('output/solar_deck.pptx'))}")
