import sys
from pathlib import Path

from pptx import Presentation


def verify(path: Path) -> None:
    if not path.is_file():
        raise SystemExit(f"Missing deck: {path}")
    deck = Presentation(path)
    if not deck.slides:
        raise SystemExit("Deck contains no slides")
    for index, slide in enumerate(deck.slides, start=1):
        text = " ".join(
            shape.text for shape in slide.shapes if hasattr(shape, "text")
        )
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        if not text.strip():
            raise SystemExit(f"Slide {index} has no text")
        if not pictures:
            raise SystemExit(f"Slide {index} has no image")
        print(f"slide {index}: text and image present")
    print(f"verified {len(deck.slides)} slides")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python inspect_deck.py output/solar_deck.pptx")
    verify(Path(sys.argv[1]))
